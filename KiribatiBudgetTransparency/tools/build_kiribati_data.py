#!/usr/bin/env python3
"""Build the data bundled into the Kiribati Budget Transparency iOS app.

Reads the Virtual Economics Kiribati hub extractions and writes one file,
KiribatiBudgetTransparency/Resources/Data/kiribati_budget.json.

Usage
  python3 tools/build_kiribati_data.py --src DIR

DIR must hold (Drive titles, as exported from the Kiribati Data Store):
  Kiribati — appropriation, supplementary and outturn 1994-2026.xlsx
  budget_<year>_summary.csv            recurrent volume front tables (Table 2 by ministry)
  budget_2023_head_tables.csv          2023 head tables (no summary file for 2023)
  budget_2026_recurrent.csv            2026 recurrent line items (no summary file for 2026)
  Kiribati — 2026 Development Budget by donor (Table 3).csv
  Kiribati — CPI by division, monthly 2006-.csv
  KNSO GDP master table 2024 — full extraction (Tables 1,2,3,6,7).csv
  Kiribati_IMF_Government_Operations_Spliced.xlsx

Every total is checked against the printed Grand Total; mismatches are
printed and carried into the app as flags, never silently corrected.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import glob
import json
import re
import sys
from collections import OrderedDict, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "KiribatiBudgetTransparency" / "Resources" / "Data" / "kiribati_budget.json"
BASE_YEAR = 2025

HEADLINE_XLSX = "Kiribati — appropriation, supplementary and outturn 1994-2026.xlsx"
DONOR_CSV = "Kiribati — 2026 Development Budget by donor (Table 3).csv"
CPI_CSV = "Kiribati — CPI by division, monthly 2006-.csv"
KNSO_GDP_CSV = "KNSO GDP master table 2024 — full extraction (Tables 1,2,3,6,7).csv"
IMF_XLSX = "Kiribati_IMF_Government_Operations_Spliced.xlsx"

# Ministries are renamed and renumbered across volumes (Health was head 22 in
# 2014 and head 14 in 2025), so they are matched on name. First match wins.
ALIASES = [
    ("beretitenti", "Office of Te Beretitenti"),
    ("public service office", "Public Service Office"),
    ("judiciary", "Judiciary"),
    ("police", "Kiribati Police Service"),
    ("public service commission", "Public Service Commission"),
    ("foreign affairs", "Ministry of Foreign Affairs and Immigration"),
    ("internal", "Ministry of Culture and Internal Affairs"),
    ("environment", "Ministry of Environment, Lands and Agricultural Development"),
    ("maneaba", "Maneaba ni Maungatabu"),
    ("commerce", "Ministry of Tourism, Commerce, Industry and Cooperatives"),
    ("audit", "Kiribati Audit Office"),
    ("attorney", "Office of the Attorney General"),
    ("fisheries", "Ministry of Fisheries and Ocean Resources"),
    ("health", "Ministry of Health and Medical Services"),
    ("education", "Ministry of Education"),
    ("communication", "Ministry of Information, Communication and Transport"),
    ("information", "Ministry of Information, Communication and Transport"),
    ("finance", "Ministry of Finance and Economic Development"),
    ("women", "Ministry of Women, Youth, Sport and Social Affairs"),
    ("public works", "Ministry of Infrastructure and Sustainable Energy"),
    ("infrastructure", "Ministry of Infrastructure and Sustainable Energy"),
    ("labour", "Ministry of Employment and Human Resources"),
    ("employment", "Ministry of Employment and Human Resources"),
    ("line and phoenix", "Ministry of Line and Phoenix Islands Development"),
    ("justice", "Ministry of Justice"),
    ("leadership", "Leadership Commission"),
    ("debt", "Debt Servicing"),
    ("subsid", "Subsidies, Grants and Other Commitments"),
    ("development fund", "Contributions to the Development Fund"),
    ("contributions to development", "Contributions to the Development Fund"),
    ("people's lawyer", "Office of the People's Lawyer"),
    ("lcdf", "Contributions to the Development Fund"),
    ("rerf", "Contributions to the RERF"),
]


def canonical(name: str) -> str:
    # Older volumes lose spaces in the text layer ("Ministryofinternal…"), so
    # match on letters only.
    compact = re.sub(r"[^a-z]", "", name.lower())
    for key, canon in ALIASES:
        if re.sub(r"[^a-z]", "", key) in compact:
            return canon
    return re.sub(r"\s+", " ", name).strip().title()


def num(s) -> float | None:
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return float(s)
    t = str(s).strip().replace(",", "").replace("$", "")
    if t in ("", "-", "()", "n.a.", "na"):
        return None
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    try:
        v = float(t)
    except ValueError:
        return None
    return -v if neg else v


# ---------------------------------------------------------------- headline

def headline(src: Path):
    import openpyxl
    wb = openpyxl.load_workbook(src / HEADLINE_XLSX, data_only=True)
    ws = wb["Appropriation and outturn"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr_i = next(i for i, r in enumerate(rows) if r and r[0] == "Year")
    keys = ["year", "appropriation", "dev_fund", "statutory", "total_operating",
            "supp_appropriated", "supp_dev_fund", "supp_statutory", "supp_total",
            "revised", "actual", "actual_dev_fund"]
    out = []
    for r in rows[hdr_i + 1:]:
        if not isinstance(r[0], int):
            break
        rec = {k: (r[i] if isinstance(r[i], (int, float)) else None) for i, k in enumerate(keys)}
        rec["year"] = int(r[0])
        out.append(rec)
    notes = [r[0] for r in rows[hdr_i + 1 + len(out):] if r and isinstance(r[0], str)]

    # "In 2025 prices" sheet: its supplementary column is the complete series
    # (workbook tension T45); carry its nominal equivalent as a separate field.
    ws2 = wb["In 2025 prices"]
    r2 = [list(r) for r in ws2.iter_rows(values_only=True)]
    h2 = next(i for i, r in enumerate(r2) if r and r[0] == "Year")
    supp_real = {}
    cpi_book = {}
    for r in r2[h2 + 1:]:
        if not isinstance(r[0], int):
            break
        if isinstance(r[5], (int, float)) and isinstance(r[10], (int, float)) and r[10]:
            supp_real[int(r[0])] = round(r[5] / r[10])  # back to nominal
        if isinstance(r[9], (int, float)):
            cpi_book[int(r[0])] = r[9]
    for rec in out:
        v = supp_real.get(rec["year"])
        if v is not None and (rec["supp_appropriated"] is None or abs(v - (rec["supp_appropriated"] or 0)) > 1):
            rec["supp_from_real_sheet"] = v

    ws3 = wb["Reconciliation"]
    r3 = [list(r) for r in ws3.iter_rows(values_only=True)]
    h3 = next(i for i, r in enumerate(r3) if r and r[0] == "Year")
    recon, cur = [], None
    for r in r3[h3 + 1:]:
        if r[0] == "Result":
            break
        if isinstance(r[0], int):
            cur = {"year": r[0], "steps": []}
            recon.append(cur)
        if cur is None or not r[1]:
            continue
        step = str(r[1])
        if step == "RESIDUAL":
            cur["residual"] = r[2]
            cur["residual_pct"] = r[3]
            cur["verdict"] = r[4]
        elif step.startswith("Revised budget reported"):
            cur["revised"] = r[3]
            cur["revised_source"] = r[4]
        else:
            cur["steps"].append({"step": step, "amount": r[2], "running": r[3], "source": r[4]})
    recon_notes = []
    tail = False
    for r in r3:
        if r and r[0] == "Result":
            tail = True
            continue
        if tail and r and isinstance(r[0], str):
            recon_notes.append(r[0])
    return out, notes, recon, recon_notes, cpi_book


# ---------------------------------------------------------------- ministries

def table2(path: Path):
    rows = [r for r in csv.DictReader(open(path, encoding="utf-8-sig")) if r["table_key"] == "appropriated_statutory"]
    labels = {}
    for r in rows:
        if r["col_label"].strip():
            labels[r["col_no"]] = r["col_label"].lower()
    net_col = next((c for c, l in labels.items() if "net provision" in l), None)
    stat_col = next((c for c, l in labels.items() if "statutory" in l), None)
    if net_col is None:  # unlabelled table (2020): the net provision is the last column
        net_col = max((r["col_no"] for r in rows), key=int)
    grouped = OrderedDict()
    for r in rows:
        k = int(r["row_no"])
        g = grouped.setdefault(k, {"label": r["row_label"].strip(), "type": r["row_type"], "page": r["page"],
                                   "source": r["source_file"], "vals": {}})
        g["vals"][r["col_no"]] = r["value"]
    lines, grand = [], None
    for g in grouped.values():
        label, vals = g["label"], g["vals"]
        if g["type"] == "total" and "grand" in label.lower():
            grand = {"operating": max(v for c, v in ((c, num(x)) for c, x in vals.items()) if v is not None),
                     "page": g["page"]}
            break
        code = vals.get("1")
        # A ministry whose name wrapped onto two lines keeps its head code but
        # loses its label; an unlabelled subtotal has neither. Keep the former.
        if g["type"] != "line" or (not label and not code):
            continue
        nums = {c: num(v) for c, v in vals.items() if c != "1" and num(v) is not None}
        if not nums:
            continue
        operating = max(nums.values())
        if operating < 10000:  # table-header fragments (e.g. a bare "2020") are not lines
            continue
        net = nums.get(net_col)
        stat = nums.get(stat_col) if stat_col else None
        if net is None:
            net = operating - stat if stat is not None else operating
        if stat is None:
            stat = operating - net
        lines.append({"label": label, "code": code, "operating": operating, "statutory": stat,
                      "net": net, "page": int(g["page"]) if str(g["page"]).isdigit() else None})
    source = rows[0]["source_file"] if rows else ""
    return lines, grand, source


def head_totals_2023(path: Path):
    """2023: sum each head's expenditure line items (no front table extracted)."""
    by_head, names, pages, stat = defaultdict(float), {}, {}, defaultdict(float)
    source = ""
    for r in csv.DictReader(open(path, encoding="utf-8-sig")):
        if r.get("section") != "EXPENDITURE" or r.get("row_type") != "line":
            continue
        v = num(r.get("budget_current"))
        if v is None:
            continue
        h = r["head_no"]
        by_head[h] += v
        names[h] = r["head_name"]
        pages.setdefault(h, r.get("page"))
        source = r.get("source_file", source)
    return [{"label": names[h], "code": h, "operating": by_head[h], "statutory": None, "net": None,
             "page": int(pages[h]) if str(pages[h]).isdigit() else None} for h in by_head], source


def head_totals_2026(path: Path):
    by_head, names, pages, stat = defaultdict(float), {}, {}, defaultdict(float)
    source, dropped = "", []
    for r in csv.DictReader(open(path, encoding="utf-8-sig")):
        if r.get("section") != "EXPENDITURE":
            continue
        v = num(r.get("b2026"))
        if v is None:
            continue
        if abs(v) > 1e10:  # extraction error (e.g. a 1.01e15 cell); dropped and reported
            dropped.append((r["head_no"], r.get("description"), v))
            continue
        h = r["head_no"]
        by_head[h] += v
        if str(r.get("statutory")) == "1":
            stat[h] += v
        names[h] = r["head_name"]
        pages.setdefault(h, r.get("page"))
        source = r.get("source_file", source)
    lines = [{"label": names[h], "code": h, "operating": by_head[h], "statutory": stat[h],
              "net": by_head[h] - stat[h], "page": int(pages[h]) if str(pages[h]).isdigit() else None}
             for h in by_head]
    return lines, source, dropped


def ministries(src: Path, head: dict[int, dict]):
    facts = defaultdict(list)          # canon -> [(year, operating, statutory, net, code, page, vol)]
    names = defaultdict(set)
    volumes, checks = {}, []

    parsed = []  # (year, lines, source, method, printed) — names resolved after all volumes are read

    def add(year, lines, source, method, printed_total):
        parsed.append((year, lines, source, method, printed_total))

    def emit(year, lines, source, method, printed_total):
        vol = len(volumes)
        volumes[str(year)] = {"file": Path(source).name if source else "", "method": method}
        total = sum(l["operating"] for l in lines)
        checks.append({"year": year, "method": method, "sum_of_lines": round(total),
                       "printed_total": printed_total,
                       "gap": round(total - printed_total) if printed_total else None})
        for l in lines:
            c = canonical(l["label"])
            names[c].add(re.sub(r"\s+", " ", l["label"]).strip())
            facts[c].append([year, round(l["operating"]),
                             None if l["statutory"] is None else round(l["statutory"]),
                             None if l["net"] is None else round(l["net"]),
                             str(l["code"] or ""), l["page"]])

    for f in sorted(glob.glob(str(src / "budget_*_summary.csv"))):
        m = re.search(r"budget_(\d{4})", f)
        year = int(m.group(1))
        lines, grand, source = table2(Path(f))
        if not lines:
            continue
        printed = grand["operating"] if grand else (head.get(year) or {}).get("total_operating")
        add(year, lines, source, "Table 2 (front table)", printed)

    # 2023 and 2026 have no front table: their head tables book the Development
    # Fund contribution, subsidies and debt servicing inside ministries, so
    # summing them would not be comparable with the Table 2 years. They are
    # reported, not charted.
    excluded = []
    p23 = src / "budget_2023_head_tables.csv"
    if p23.exists():
        lines, source = head_totals_2023(p23)
        excluded.append({"year": 2023, "sum_of_lines": round(sum(l["operating"] for l in lines)),
                         "printed_total": (head.get(2023) or {}).get("total_operating"),
                         "reason": "no Table 2 extracted; head tables include the Development Fund, subsidies and debt servicing within ministries"})
    p26 = src / "budget_2026_recurrent.csv"
    if p26.exists():
        lines, source, dropped = head_totals_2026(p26)
        for d in dropped:
            print(f"  2026: dropped implausible cell head {d[0]} {d[1]!r} = {d[2]:.3g}")
        excluded.append({"year": 2026, "sum_of_lines": round(sum(l["operating"] for l in lines)),
                         "printed_total": (head.get(2026) or {}).get("total_operating"),
                         "reason": "no Table 2 extracted; line items include the Development Fund, subsidies and debt servicing within ministries"})

    # Name unlabelled lines from the same head code in the nearest year's volume.
    code_names = defaultdict(dict)
    for year, lines, *_ in parsed:
        for l in lines:
            if l["label"] and l["code"]:
                code_names[str(l["code"])][year] = l["label"]
    for year, lines, source, method, printed in sorted(parsed, key=lambda t: t[0]):
        for l in lines:
            if not l["label"]:
                cands = code_names.get(str(l["code"]), {})
                if cands:
                    near = min(cands, key=lambda y: abs(y - year))
                    l["label"] = cands[near]
                    l["label_inferred"] = True
                    print(f"  {year}: unlabelled head {l['code']} named from {near}: {cands[near]!r}")
                else:
                    l["label"] = f"Head {l['code']}"
        emit(year, lines, source, method, printed)

    out = []
    for c, fs in facts.items():
        # merge duplicate lines for the same year (e.g. two rows mapped to one ministry)
        by_year = OrderedDict()
        for f in sorted(fs):
            if f[0] in by_year:
                p = by_year[f[0]]
                p[1] += f[1]
                p[2] = None if p[2] is None or f[2] is None else p[2] + f[2]
                p[3] = None if p[3] is None or f[3] is None else p[3] + f[3]
            else:
                by_year[f[0]] = list(f)
        out.append({"name": c, "printed_names": sorted(names[c]),
                    "facts": [{"year": f[0], "operating": f[1], "statutory": f[2], "net": f[3],
                               "code": f[4], "page": f[5]} for f in by_year.values()]})
    out.sort(key=lambda m: -max(f["operating"] for f in m["facts"]))
    return out, volumes, checks, excluded


# ---------------------------------------------------------------- donors

def donors(src: Path):
    p = src / DONOR_CSV
    if not p.exists():
        return None
    rows = list(csv.DictReader(open(p, encoding="utf-8-sig")))
    cols = list(rows[0].keys())
    out = [{"donor": r[cols[0]].strip(), "values": [num(r[c]) for c in cols[1:]]} for r in rows]
    return {"columns": cols[1:], "rows": out, "source": DONOR_CSV}


# ---------------------------------------------------------------- denominators

def denominators(src: Path, head: dict[int, dict]):
    # CPI: annual average of the monthly all-items index (PDH column after Sep 2024).
    months = defaultdict(list)
    for r in csv.DictReader(open(src / CPI_CSV, encoding="utf-8-sig")):
        y = int(r["month"][:4])
        v = num(r.get("All-Items CPI")) or num(r.get("PDH DF_CPI total"))
        if v is not None:
            months[y].append(v)
    cpi_avg = {y: sum(v) / len(v) for y, v in months.items() if len(v) >= 9}
    base = cpi_avg[BASE_YEAR]
    cpi = {y: v / base * 100 for y, v in cpi_avg.items()}
    cpi_months = {y: len(months[y]) for y in cpi}

    # GDP: KNSO Table 1 (A$'000) to 2021, IMF (A$m) from 2022.
    knso = {}
    for r in csv.DictReader(open(src / KNSO_GDP_CSV, encoding="utf-8-sig")):
        if r["sheet"] == "Table 1" and r["series"].strip() == "Gross domestic product (GDP) in current prices":
            v = num(r["value"])
            if v is not None and v > 1000:     # levels only; growth-rate rows share the label
                knso.setdefault(int(r["year"]), v / 1000)
    import openpyxl
    wb = openpyxl.load_workbook(src / IMF_XLSX, read_only=True, data_only=True)
    rows = [list(r) for r in wb["Master Spliced"].iter_rows(values_only=True)]
    hdr = next(r for r in rows if r and r[0] == "Line item")
    ycol = {c: v for c, v in enumerate(hdr) if isinstance(v, int)}

    def imf_row(label):
        r = next(r for r in rows if r and r[0] == label)
        return {ycol[c]: float(r[c]) for c in ycol if isinstance(r[c], (int, float))}
    imf_gdp = imf_row("Nominal GDP")
    imf_exp = imf_row("Total expenditure")
    obs = [list(r) for r in wb["Observation Map"].iter_rows(values_only=True)]
    status = {}
    for r in obs:
        if r and r[0] in ("Nominal GDP", "Total expenditure") and isinstance(r[1], int):
            status[(r[0], r[1])] = {"Historical/actual": "a", "Estimate": "e", "Projection": "p"}.get(str(r[5]), "a")

    out = {}
    for y in range(2005, 2032):
        rec = {}
        if y in knso:
            rec.update(gdp=round(knso[y], 1), gdp_src="knso", gdp_status="a")
        elif y in imf_gdp:
            rec.update(gdp=imf_gdp[y], gdp_src="imf", gdp_status=status.get(("Nominal GDP", y), "a"))
        if y in cpi:
            rec.update(cpi=round(cpi[y], 3), cpi_months=cpi_months[y])
        if y in imf_exp:
            rec.update(imf_total_exp=imf_exp[y], imf_exp_status=status.get(("Total expenditure", y), "a"))
        if rec:
            out[str(y)] = rec
    return out


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--src", type=Path, required=True)
    a = ap.parse_args()
    src = a.src

    head, notes, recon, recon_notes, cpi_book = headline(src)
    by_year = {h["year"]: h for h in head}
    mins, volumes, checks, excluded = ministries(src, by_year)
    dens = denominators(src, by_year)

    for e in excluded:
        print(f"  {e['year']}: not charted ({e['reason']}); lines {e['sum_of_lines']:,} vs printed {e['printed_total']:,}")
    for c in checks:
        flag = "" if not c["gap"] else f"   <-- gap {c['gap']:+,}"
        print(f"  {c['year']}: lines {c['sum_of_lines']:>13,}  printed {c['printed_total'] or 0:>13,}  ({c['method']}){flag}")
    for y, v in cpi_book.items():
        mine = dens.get(str(y), {}).get("cpi")
        if mine and abs(mine / 100 * dens[str(BASE_YEAR)]["cpi"] - mine) > 1e9:
            pass

    payload = {
        "generated": dt.date.today().isoformat(),
        "base_year": BASE_YEAR,
        "currency": "AUD",
        "headline": head,
        "headline_notes": notes,
        "reconciliation": recon,
        "reconciliation_notes": recon_notes,
        "ministries": mins,
        "ministry_volumes": volumes,
        "ministry_checks": checks,
        "ministry_years_not_charted": excluded,
        "donors_2026": donors(src),
        "denominators": dens,
        "sources": {
            "headline": HEADLINE_XLSX,
            "cpi": "Kiribati National Statistics Office all-items CPI, annual average of the monthly index "
                   "(KNSO release to Sep 2024; Pacific Data Hub DF_CPI from Oct 2024)",
            "gdp": "KNSO national accounts, GDP master table 2024, Table 1 (to 2021); "
                   "IMF Country Report 26/099 central government operations, nominal GDP memo line (2022 onward)",
            "total_exp": "IMF central government operations, total expenditure (spliced 2009–2031)",
            "ministries": "MFED recurrent budget volumes, Table 2 — appropriated and statutory expenditure by ministry",
            "donors": DONOR_CSV,
        },
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"wrote {OUT.relative_to(HERE.parent)} ({OUT.stat().st_size/1024:,.0f} KB): "
          f"{len(head)} headline years, {len(mins)} ministries, {len(dens)} denominator years")


if __name__ == "__main__":
    main()
