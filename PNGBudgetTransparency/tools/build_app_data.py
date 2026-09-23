#!/usr/bin/env python3
"""Refresh the data bundled into the PNG Budget Transparency iOS app.

The app reads the same payload the Virtual Economics data site publishes at
/png/budget/ (the <script id="payload"> blob), so the site and the app can never
disagree about a figure. Optionally it also bundles the Public Investment
Programme projects from the site's pip_projects.parquet export.

Usage
  python3 tools/build_app_data.py --site path/to/png-budget-site/png/budget/index.html \
      [--pip path/to/png-budget-site/png/budget/data/pip_projects.parquet]

  # or, if you already have the payload as JSON:
  python3 tools/build_app_data.py --payload budget_payload.json

Writes into PNGBudgetTransparency/Resources/Data/:
  budget_payload.json   minified site payload (required by the app)
  pip_projects.json     compact PIP table (optional; the Projects tab hides without it)
  anu_denominators.json GDP, prices and total spending from the ANU PNG National
                        Budget Database (drives the constant-price, % of GDP and
                        % of spending views)

  python3 tools/build_app_data.py --site ... \
      --anu-budget "2025-12 — PNG National Budget Database — 2026 Budget.xlsx" \
      --anu-econ   "2026-08-12 — PNG Economic Database — Dashboard1 crosstab (Maindatabase30June).csv"
"""
from __future__ import annotations

import argparse
import html
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "PNGBudgetTransparency" / "Resources" / "Data"

REQUIRED_KEYS = {"generated", "corpus_version", "measures", "fiscal", "agencies",
                 "series_codes", "volumes", "facts", "counts"}


def payload_from_site(index_html: Path) -> dict:
    text = index_html.read_text(encoding="utf-8")
    m = re.search(r'<script id="payload" type="application/json">(.*?)</script>', text, re.S)
    if not m:
        sys.exit(f"{index_html}: no <script id=\"payload\"> block found")
    raw = m.group(1).strip()
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        return json.loads(html.unescape(raw))


def check_payload(p: dict) -> None:
    missing = REQUIRED_KEYS - p.keys()
    if missing:
        sys.exit(f"payload is missing keys the app needs: {sorted(missing)}")
    n_ag = len(p["agencies"])
    for i, f in enumerate(p["facts"]):
        if len(f) != 8:
            sys.exit(f"facts[{i}] has {len(f)} fields, expected 8 "
                     "[agencyIdx, edition, refYear, serIdx, value, printedPage, pdfPage, volIdx]")
        if not (0 <= f[0] < n_ag and 0 <= f[3] < len(p["series_codes"]) and 0 <= f[7] < len(p["volumes"])):
            sys.exit(f"facts[{i}] points outside agencies/series_codes/volumes: {f}")


def pip_table(parquet: Path) -> dict:
    """Compact the PIP export into the shape PIPData.swift decodes.

    projects: one per PIP number, named as its most recent edition names it
    facts:    [projectIdx, edition, refYear, value K m, printedPage, pdfPage, volumeIdx]
              refYear 0 marks the five-year total as the volume prints it
    """
    try:
        import pyarrow.parquet as pq
    except ImportError:
        sys.exit("pip install pyarrow to convert the PIP parquet export")
    rows = pq.read_table(parquet).to_pylist()
    need = {"budget_edition", "executing_agency_code", "executing_agency_name", "programme_group",
            "pip_number", "project_name", "ref_year", "value_kina_m", "source_volume",
            "printed_page", "pdf_page"}
    missing = need - set(rows[0]) if rows else need
    if missing:
        sys.exit(f"{parquet}: missing columns {sorted(missing)}")

    latest: dict[str, dict] = {}
    for r in rows:
        k = r["pip_number"]
        if k not in latest or r["budget_edition"] >= latest[k]["budget_edition"]:
            latest[k] = r
    numbers = sorted(latest)
    pidx = {n: i for i, n in enumerate(numbers)}
    volumes = sorted({r["source_volume"] for r in rows})
    vidx = {v: i for i, v in enumerate(volumes)}
    projects = [{
        "pip": n,
        "name": latest[n]["project_name"].strip(),
        "agency_code": latest[n]["executing_agency_code"],
        "agency": latest[n]["executing_agency_name"].strip(),
        "group": latest[n]["programme_group"],
        "names": sorted({r["project_name"].strip() for r in rows if r["pip_number"] == n} - {latest[n]["project_name"].strip()}),
    } for n in numbers]
    facts = [[pidx[r["pip_number"]], r["budget_edition"], (0 if r["ref_year"] == "5yr_total" else int(r["ref_year"])), round(r["value_kina_m"], 3),
              r["printed_page"], r["pdf_page"], vidx[r["source_volume"]]] for r in rows]
    facts.sort(key=lambda f: (f[0], f[1], f[2]))
    return {"volumes": volumes, "projects": projects, "facts": facts}


BASE_YEAR = 2025


def anu_denominators(budget_xlsx: Path, econ_csv: Path | None) -> dict:
    """Denominators from the ANU Development Policy Centre / UPNG PNG National
    Budget Database, read by row label and aligned by column so years cannot slip.

    gdp        Analysis sheet "Nominal GDP" (Treasury/NSO new series from 2007);
               earlier years use the ANU PNG Economic Database's
               "GDP (current prices, new series)" so old and new series are never mixed
    cpi        index, BASE_YEAR = 100, chained from Table 9 "Average on Average (%)"
    deflator   index, BASE_YEAR = 100, Analysis "Deflator base year 2022" rebased
    total_exp  Analysis "Total expenditure & net lending"
    """
    try:
        import openpyxl
    except ImportError:
        sys.exit("pip install openpyxl to read the ANU budget database")
    wb = openpyxl.load_workbook(budget_xlsx, read_only=True, data_only=True)

    def grid(sheet):
        return [list(r) for r in wb[sheet].iter_rows(values_only=True)]

    def year_header(rows, max_scan=40):
        for r in rows[:max_scan]:
            cols = {c: v for c, v in enumerate(r) if isinstance(v, int) and 1900 < v < 2100}
            if len(cols) > 10:
                return cols
        sys.exit(f"{budget_xlsx}: no year header found")

    def row_by_label(rows, ycols, label, first_only=True):
        first_data_col = min(ycols)
        for r in rows:
            labels = [str(x).strip() for x in r[:first_data_col] if x is not None]
            if label in labels:
                return {ycols[c]: float(r[c]) for c in ycols if c < len(r) and isinstance(r[c], (int, float))}
        sys.exit(f"{budget_xlsx}: row '{label}' not found")

    an = grid("Analysis")
    ay = year_header(an)
    gdp = row_by_label(an, ay, "Nominal GDP")
    d22 = row_by_label(an, ay, "Deflator base year 2022")
    exp = row_by_label(an, ay, "Total expenditure & net lending")

    t9 = grid("Prices (Tb9)")
    ty = year_header(t9)
    infl = row_by_label(t9, ty, "Average on Average (%)")          # first block = latest budget vintage
    t9_status = {ty[c]: str(t9[1][c]).strip().lower() for c in ty if c < len(t9[1]) and t9[1][c]}

    t1 = grid("GDP (Tb1)")
    gy = year_header(t1)
    gdp_status = {gy[c]: str(t1[1][c]).strip().lower() for c in gy if c < len(t1[1]) and t1[1][c]}

    # CPI index chained from annual-average inflation, BASE_YEAR = 100.
    years = sorted(infl)
    cpi = {years[0]: 100.0}
    for a, b in zip(years, years[1:]):
        cpi[b] = cpi[a] * (1 + infl[b] / 100)
    base = cpi[BASE_YEAR]
    cpi = {y: v / base * 100 for y, v in cpi.items()}
    deflator = {y: v / d22[BASE_YEAR] * 100 for y, v in d22.items()}

    econ_gdp: dict[int, float] = {}
    if econ_csv:
        import csv
        with open(econ_csv, encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                if r["Variable"] == "GDP (current prices, new series)" and r["Amount"].strip():
                    econ_gdp[int(r["Year"])] = float(r["Amount"].replace(",", ""))

    def status(tag: str | None) -> str:
        return {"actual": "a", "estimate": "e", "projection": "p"}.get((tag or "").lower(), "a")

    NEW_SERIES_FROM = 2007
    out = {}
    for y in range(2000, max(ay.values()) + 1):
        rec = {}
        if y >= NEW_SERIES_FROM and y in gdp:
            rec["gdp"] = round(gdp[y], 1)
            rec["gdp_src"] = "b"
            rec["gdp_status"] = status(gdp_status.get(y))
        elif y in econ_gdp:
            rec["gdp"] = round(econ_gdp[y], 1)
            rec["gdp_src"] = "e"
            rec["gdp_status"] = "a"
        if y in cpi:
            rec["cpi"] = round(cpi[y], 3)
            rec["cpi_status"] = status(t9_status.get(y))
        if y in deflator:
            rec["deflator"] = round(deflator[y], 3)
        if y in exp:
            rec["total_exp"] = round(exp[y], 1)
            rec["exp_status"] = "a" if y <= 2024 else ("e" if y == 2025 else "p")
        if rec:
            out[str(y)] = rec

    notes = [r[1] for r in grid("Notes")[:3] if len(r) > 11]
    updated = wb["Notes"].cell(row=2, column=12).value
    return {
        "base_year": BASE_YEAR,
        "source": {
            "name": "PNG National Budget Database",
            "publisher": "ANU Development Policy Centre and University of Papua New Guinea",
            "edition": str(wb["Notes"].cell(row=3, column=12).value or "").strip(),
            "authors": str(wb["Notes"].cell(row=4, column=12).value or "").strip(),
            "updated": updated.date().isoformat() if hasattr(updated, "date") else str(updated or ""),
            "file": budget_xlsx.name,
            "econ_file": econ_csv.name if econ_csv else None,
        },
        "years": out,
    }


def write_json(obj, path: Path) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"wrote {path.relative_to(HERE.parent)}  ({path.stat().st_size / 1024:,.0f} KB)")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--site", type=Path, help="the site's png/budget/index.html")
    src.add_argument("--payload", type=Path, help="payload JSON already extracted")
    ap.add_argument("--pip", type=Path, help="pip_projects.parquet from the site's data/ folder")
    ap.add_argument("--anu-budget", type=Path, help="ANU PNG National Budget Database workbook (.xlsx)")
    ap.add_argument("--anu-econ", type=Path, help="ANU PNG Economic Database crosstab (.csv), for pre-2007 GDP")
    a = ap.parse_args()

    payload = payload_from_site(a.site) if a.site else json.loads(a.payload.read_text(encoding="utf-8"))
    check_payload(payload)
    OUT.mkdir(parents=True, exist_ok=True)
    write_json(payload, OUT / "budget_payload.json")
    print(f"  generated {payload['generated']}, corpus {payload['corpus_version']}, "
          f"{len(payload['fiscal'])} fiscal records, {len(payload['agencies'])} agencies, "
          f"{len(payload['facts'])} agency facts")

    if a.pip:
        write_json(pip_table(a.pip), OUT / "pip_projects.json")

    if a.anu_budget:
        anu = anu_denominators(a.anu_budget, a.anu_econ)
        write_json(anu, OUT / "anu_denominators.json")
        ys = sorted(int(y) for y in anu["years"])
        print(f"  ANU {anu['source']['edition']} (updated {anu['source']['updated']}), years {ys[0]}–{ys[-1]}")


if __name__ == "__main__":
    main()
